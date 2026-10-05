from __future__ import annotations

import html
import json
import os
import posixpath
import re
import shutil
import socket
import subprocess
import sys
import time
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright


BASE = Path(__file__).resolve().parent
DRIVE_FOLDER_ID = "1TTdT09m-WOVeEiii2hed9y4JhsNw8uh9"
URL_DRIVE = f"https://drive.google.com/drive/u/0/folders/{DRIVE_FOLDER_ID}"

PERFIL_GOOGLE_EDGE = Path(r"C:\HAGAP\GOOGLE_DRIVE_BROWSER_PROFILE")
CDP_GOOGLE_PORT = 9223

RE_PROJ = re.compile(r"(?<!\d)(\d{7})(?!\d)")

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def escolher_planilha() -> Path:
    candidatos = [
        Path.cwd() / "PROGRAMAÇÃO UMU .xlsx",
        BASE / "PROGRAMAÇÃO UMU .xlsx",
        Path.home() / "Downloads" / "PROGRAMAÇÃO UMU .xlsx",
        Path.home() / "Desktop" / "PROGRAMAÇÃO UMU .xlsx",
        Path.home() / "Documents" / "PROGRAMAÇÃO UMU .xlsx",
    ]
    for p in candidatos:
        if p.exists():
            return p

    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        arq = filedialog.askopenfilename(
            title="Selecione PROGRAMAÇÃO UMU .xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        root.destroy()
        if arq:
            return Path(arq)
    except Exception:
        pass

    raise RuntimeError("PROGRAMAÇÃO UMU .xlsx não localizada.")


def encontrar_edge() -> Path:
    candidatos = [
        Path(os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)")) / "Microsoft/Edge/Application/msedge.exe",
        Path(os.environ.get("PROGRAMFILES", r"C:\Program Files")) / "Microsoft/Edge/Application/msedge.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft/Edge/Application/msedge.exe",
    ]
    for p in candidatos:
        if p.exists():
            return p
    raise RuntimeError("Microsoft Edge não localizado.")


def porta_aberta(host: str, port: int) -> bool:
    try:
        with socket.create_connection((host, port), timeout=0.5):
            return True
    except OSError:
        return False


def abrir_edge_google_normal():
    if porta_aberta("127.0.0.1", CDP_GOOGLE_PORT):
        return

    PERFIL_GOOGLE_EDGE.mkdir(parents=True, exist_ok=True)
    edge = encontrar_edge()

    subprocess.Popen(
        [
            str(edge),
            f"--remote-debugging-port={CDP_GOOGLE_PORT}",
            "--remote-allow-origins=*",
            f"--user-data-dir={PERFIL_GOOGLE_EDGE}",
            "--no-first-run",
            "--no-default-browser-check",
            "--new-window",
            URL_DRIVE,
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        creationflags=getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0),
    )

    inicio = time.time()
    while time.time() - inicio < 30:
        if porta_aberta("127.0.0.1", CDP_GOOGLE_PORT):
            return
        time.sleep(0.5)

    raise RuntimeError("O Edge abriu, mas não consegui conectar ao navegador.")


def conectar_drive(playwright):
    abrir_edge_google_normal()
    browser = playwright.chromium.connect_over_cdp(
        f"http://127.0.0.1:{CDP_GOOGLE_PORT}",
        timeout=30000,
    )
    contexts = browser.contexts
    if not contexts:
        raise RuntimeError("Edge do Google abriu sem contexto navegável.")
    context = contexts[0]

    for page in context.pages:
        if "drive.google.com" in (page.url or "").lower():
            return browser, page

    page = context.new_page()
    page.goto(URL_DRIVE, wait_until="domcontentloaded", timeout=120000)
    return browser, page


def drive_autenticado(page) -> bool:
    try:
        url = (page.url or "").lower()
        titulo = (page.title() or "").lower()
        if "accounts.google.com" in url:
            return False
        if "drive.google.com" not in url:
            return False
        if "sign in" in titulo or "fazer login" in titulo:
            return False

        # Sinais reais da interface do Drive.
        sinais = 0
        for seletor in [
            'button[guidedhelpid="new_menu_button"]',
            'div[role="main"]',
            'tr[data-id]',
            '[data-target="navTree"]',
        ]:
            try:
                if page.locator(seletor).count():
                    sinais += 1
            except Exception:
                pass
        return sinais >= 2
    except Exception:
        return False


def aguardar_drive(page, limite=600):
    page.goto(URL_DRIVE, wait_until="domcontentloaded", timeout=120000)
    inicio = time.time()
    while time.time() - inicio < limite:
        if drive_autenticado(page):
            if DRIVE_FOLDER_ID not in (page.url or ""):
                page.goto(URL_DRIVE, wait_until="domcontentloaded", timeout=120000)
                time.sleep(2)
            if drive_autenticado(page):
                print("[CONFIRMADO] Google Drive autenticado e pasta 'projetos' aberta.")
                return
        print(
            "\r[PENDENTE] Faça login no Google e deixe a pasta 'projetos' aberta...       ",
            end="",
            flush=True,
        )
        time.sleep(2)
    print()
    raise RuntimeError("Google Drive não ficou autenticado dentro do tempo limite.")


def coletar_itens_drive(page):
    """
    Percorre a lista virtualizada da pasta 'projetos'.
    Coleta ID real + nome. O link universal open?id= funciona para arquivo e pasta.
    """
    page.bring_to_front()
    page.goto(URL_DRIVE, wait_until="domcontentloaded", timeout=120000)
    time.sleep(3)

    itens = {}
    sem_novos = 0
    ultima_qtd = 0

    for rodada in range(700):
        lote = page.evaluate(
            """
            () => {
                const out = [];
                for (const row of document.querySelectorAll('tr[data-id]')) {
                    const id = row.getAttribute('data-id');
                    if (!id) continue;

                    let nome = '';
                    const forte = row.querySelector('strong.DNoYtb');
                    if (forte) nome = (forte.innerText || '').trim();

                    if (!nome) {
                        const titulo = row.querySelector('[data-id="' + id + '"][data-tooltip]');
                        if (titulo) nome = (titulo.getAttribute('data-tooltip') || '').trim();
                    }

                    if (!nome) {
                        const linhas = (row.innerText || '').split(String.fromCharCode(10)).map(x => x.trim()).filter(Boolean);
                        if (linhas.length) nome = linhas[0];
                    }

                    if (nome) out.push({id, nome});
                }
                return out;
            }
            """
        )

        for item in lote:
            itens[item["id"]] = item["nome"]

        if len(itens) == ultima_qtd:
            sem_novos += 1
        else:
            sem_novos = 0
            ultima_qtd = len(itens)

        print(
            f"\r[VARREDURA DRIVE] Itens coletados: {len(itens)} | rodada {rodada + 1}       ",
            end="",
            flush=True,
        )

        status = page.evaluate(
            """
            () => {
                const row = document.querySelector('tr[data-id]');
                if (!row) return {ok:false};

                let el = row.parentElement;
                while (el) {
                    const cs = getComputedStyle(el);
                    if (
                        el.scrollHeight > el.clientHeight + 50 &&
                        (cs.overflowY === 'auto' || cs.overflowY === 'scroll')
                    ) {
                        const antes = el.scrollTop;
                        el.scrollTop = Math.min(
                            el.scrollTop + Math.max(400, el.clientHeight * 0.82),
                            el.scrollHeight
                        );
                        el.dispatchEvent(new Event('scroll', {bubbles:true}));
                        return {
                            ok:true,
                            before:antes,
                            top:el.scrollTop,
                            max:el.scrollHeight,
                            client:el.clientHeight
                        };
                    }
                    el = el.parentElement;
                }
                return {ok:false};
            }
            """
        )

        if not status.get("ok"):
            # Fallback por teclado.
            try:
                page.keyboard.press("PageDown")
            except Exception:
                pass

        time.sleep(0.18)

        if status.get("ok"):
            no_fim = status.get("top", 0) + status.get("client", 0) >= status.get("max", 0) - 5
            if no_fim and sem_novos >= 4:
                break
        elif sem_novos >= 12:
            break

    print()
    return itens


def montar_mapa_projetos(itens):
    """
    Prioridade de confiabilidade:
    - número de 7 dígitos presente no NOME do item no Drive.
    - uma única ocorrência de destino por projeto.
    Divergência => não cria link automaticamente.
    """
    bruto = defaultdict(list)

    for item_id, nome in itens.items():
        projetos = set(RE_PROJ.findall(nome or ""))
        for projeto in projetos:
            bruto[projeto].append(
                {
                    "id": item_id,
                    "nome": nome,
                    "url": f"https://drive.google.com/open?id={item_id}",
                }
            )

    mapa = {}
    divergencias = {}

    for projeto, hits in bruto.items():
        unicos = {}
        for hit in hits:
            unicos[hit["id"]] = hit
        hits = list(unicos.values())

        if len(hits) == 1:
            mapa[projeto] = hits[0]["url"]
        elif len(hits) > 1:
            divergencias[projeto] = hits

    return mapa, divergencias


def projeto_da_celula(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int):
        s = str(v)
    elif isinstance(v, float) and v.is_integer():
        s = str(int(v))
    elif isinstance(v, str):
        s = v.strip()
    else:
        return None
    return s if re.fullmatch(r"\d{7}", s) else None


def mapear_celulas(planilha: Path, links):
    alvo = defaultdict(list)
    todos = 0
    wb = load_workbook(planilha, read_only=True, data_only=False, keep_links=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                p = projeto_da_celula(cell.value)
                if not p:
                    continue
                todos += 1
                if p in links:
                    alvo[ws.title].append((cell.coordinate, p, links[p]))
    abas = list(wb.sheetnames)
    wb.close()
    return alvo, todos, abas


def localizar_xml_abas(planilha: Path):
    with zipfile.ZipFile(planilha, "r") as z:
        wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
        rel_xml = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))

    rel_targets = {}
    for rel in rel_xml.findall(f"{{{NS_PKG}}}Relationship"):
        rel_targets[rel.attrib["Id"]] = rel.attrib["Target"]

    paths = {}
    sheets_node = wb_xml.find(f"{{{NS_MAIN}}}sheets")
    for sh in sheets_node.findall(f"{{{NS_MAIN}}}sheet"):
        nome = sh.attrib["name"]
        rid = sh.attrib[f"{{{NS_R}}}id"]
        target = rel_targets[rid]
        path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        paths[nome] = path
    return paths


def esc(s):
    return html.escape(str(s), quote=True)


def ensure_r_namespace(xml_text):
    if 'xmlns:r="' in xml_text:
        return xml_text
    return xml_text.replace(
        "<worksheet ",
        f'<worksheet xmlns:r="{NS_R}" ',
        1,
    )


def inserir_bloco_hyperlinks(xml_text, bloco):
    # Mantém a ordem esperada do schema do worksheet.
    tags_posteriores = [
        "printOptions", "pageMargins", "pageSetup", "headerFooter",
        "rowBreaks", "colBreaks", "customProperties", "cellWatches",
        "ignoredErrors", "smartTags", "drawing", "legacyDrawing",
        "legacyDrawingHF", "picture", "oleObjects", "controls",
        "webPublishItems", "tableParts", "extLst",
    ]
    posicoes = []
    for tag in tags_posteriores:
        m = re.search(rf"<{tag}(?:\s|/|>)", xml_text)
        if m:
            posicoes.append(m.start())
    pos = min(posicoes) if posicoes else xml_text.rfind("</worksheet>")
    return xml_text[:pos] + bloco + xml_text[pos:]


def gerar_excel_linkado(planilha: Path, mapa_links, divergencias):
    targets, total_cells, abas = mapear_celulas(planilha, mapa_links)
    sheet_paths = localizar_xml_abas(planilha)

    out = planilha.with_name(planilha.stem + " - LINKS DRIVE.xlsx")
    modified = {}

    ocorrencias = 0
    projetos_linkados = set()
    links_antigos = 0

    with zipfile.ZipFile(planilha, "r") as zin:
        nomes_original = set(zin.namelist())

        for sheet_name, items in targets.items():
            if not items:
                continue

            sheet_path = sheet_paths[sheet_name]
            xml_text = zin.read(sheet_path).decode("utf-8")
            xml_text = ensure_r_namespace(xml_text)

            rel_path = posixpath.join(
                posixpath.dirname(sheet_path),
                "_rels",
                posixpath.basename(sheet_path) + ".rels",
            )

            if rel_path in nomes_original:
                rel_text = zin.read(rel_path).decode("utf-8")
            else:
                rel_text = (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    f'<Relationships xmlns="{NS_PKG}"></Relationships>'
                )

            novos_links = []
            novas_rels = []
            contador = 1

            for coord, projeto, url in items:
                # Se já havia hyperlink nesta célula, preserva e não sobrescreve.
                if re.search(rf'<hyperlink\b[^>]*\bref="{re.escape(coord)}"[^>]*/?>', xml_text):
                    links_antigos += 1
                    continue

                while True:
                    rid = f"rIdHAGAP{contador}"
                    contador += 1
                    if f'Id="{rid}"' not in rel_text:
                        break

                novos_links.append(f'<hyperlink ref="{esc(coord)}" r:id="{rid}"/>')
                novas_rels.append(
                    f'<Relationship Id="{rid}" '
                    f'Type="{NS_R}/hyperlink" '
                    f'Target="{esc(url)}" TargetMode="External"/>'
                )
                ocorrencias += 1
                projetos_linkados.add(projeto)

            if not novos_links:
                continue

            rel_text = rel_text.replace(
                "</Relationships>",
                "".join(novas_rels) + "</Relationships>",
            )

            if "</hyperlinks>" in xml_text:
                xml_text = xml_text.replace(
                    "</hyperlinks>",
                    "".join(novos_links) + "</hyperlinks>",
                    1,
                )
            else:
                bloco = "<hyperlinks>" + "".join(novos_links) + "</hyperlinks>"
                xml_text = inserir_bloco_hyperlinks(xml_text, bloco)

            modified[sheet_path] = xml_text.encode("utf-8")
            modified[rel_path] = rel_text.encode("utf-8")

        # Copia o XLSX preservando literalmente tudo que não foi modificado.
        with zipfile.ZipFile(out, "w") as zout:
            for info in zin.infolist():
                data = modified.get(info.filename, zin.read(info.filename))
                zout.writestr(info, data)

            for name, data in modified.items():
                if name not in nomes_original:
                    zout.writestr(name, data)

    # Validação sem salvar novamente pelo openpyxl.
    original = load_workbook(planilha, read_only=True, data_only=False, keep_links=True)
    novo = load_workbook(out, read_only=True, data_only=False, keep_links=True)

    if original.sheetnames != novo.sheetnames:
        raise RuntimeError("Validação falhou: nomes/ordem das abas mudou.")

    for nome in original.sheetnames:
        wo = original[nome]
        wn = novo[nome]
        if wo.max_row != wn.max_row or wo.max_column != wn.max_column:
            raise RuntimeError(f"Validação falhou: dimensão da aba {nome} mudou.")

        for ro, rn in zip(
            wo.iter_rows(values_only=True),
            wn.iter_rows(values_only=True),
        ):
            if ro != rn:
                raise RuntimeError(f"Validação falhou: algum valor da aba {nome} mudou.")

    original.close()
    novo.close()

    relatorio = {
        "gerado_em": datetime.now().isoformat(timespec="seconds"),
        "arquivo_origem": str(planilha),
        "arquivo_saida": str(out),
        "abas_preservadas": abas,
        "celulas_com_projeto_7_digitos": total_cells,
        "projetos_com_link_drive": len(projetos_linkados),
        "ocorrencias_linkadas": ocorrencias,
        "hyperlinks_preexistentes_preservados": links_antigos,
        "divergencias_drive": divergencias,
    }

    rel_path = out.with_name(out.stem + " - RELATORIO.json")
    rel_path.write_text(json.dumps(relatorio, ensure_ascii=False, indent=2), encoding="utf-8")

    return out, rel_path, relatorio


def main():
    planilha = escolher_planilha()

    print("=" * 78)
    print("HAGAP - GERAR PROGRAMAÇÃO IDÊNTICA COM LINKS DO GOOGLE DRIVE")
    print(f"Origem: {planilha}")
    print("[CONFIRMADO] A programação original NÃO será alterada.")
    print("[CONFIRMADO] O arquivo de saída mantém valores/layout; só adiciona hyperlinks.")
    print("=" * 78)

    with sync_playwright() as p:
        browser, page = conectar_drive(p)
        aguardar_drive(page)

        itens = coletar_itens_drive(page)
        if not itens:
            raise RuntimeError("Nenhum item foi coletado da pasta projetos do Drive.")

        mapa, divergencias = montar_mapa_projetos(itens)

    print(f"[CONFIRMADO] Itens coletados no Drive: {len(itens)}")
    print(f"[CONFIRMADO] Projetos com destino único: {len(mapa)}")
    print(f"[DIVERGÊNCIA] Projetos com mais de um destino: {len(divergencias)}")

    out, rel, info = gerar_excel_linkado(planilha, mapa, divergencias)

    print()
    print("=" * 78)
    print("RESULTADO")
    print(f"Arquivo: {out}")
    print(f"Ocorrências linkadas: {info['ocorrencias_linkadas']}")
    print(f"Projetos únicos linkados: {info['projetos_com_link_drive']}")
    print(f"Relatório: {rel}")
    print("[CONFIRMADO] Valores, abas e dimensões foram reconferidos sem alteração.")
    print("=" * 78)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nCancelado.")
    except Exception as exc:
        print(f"\n[ERRO] {exc}")
        sys.exit(1)
